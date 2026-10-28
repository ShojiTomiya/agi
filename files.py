"""Attached files: the short description that goes into a message, and the text
units (pages, slides, cells, lines) that the read_file tool pages through.

A document is never pasted into the conversation whole unless it's short: the
message gets its beginning and the model reads the rest with read_file, so a long
PDF costs only what's actually read and stays readable to the end.
"""
from functools import lru_cache
from pathlib import Path

MAX_PREVIEW_ROWS = 5
# a document this short goes into the message whole
SMALL_FILE_CHARS = 3000
# how much of a longer one the message shows
HEAD_CHARS = 1500

SUPPORTED_EXTENSIONS = {
    "image": ["png", "jpg", "jpeg", "webp", "gif"],
    "pdf":   ["pdf"],
    "docx":  ["docx"],
    "notebook": ["ipynb"],
    "excel": ["xlsx", "xls"],
    "csv":   ["csv"],
    "pptx":  ["pptx"],
    "text":  ["md", "txt", "py", "js", "ts", "json", "yaml", "yml", "html", "css"],
}

ALL_SUPPORTED = [ext for exts in SUPPORTED_EXTENSIONS.values() for ext in exts]

UNIT = {"pdf": "page", "pptx": "slide", "notebook": "cell"}   # everything else: line
LABEL = {"pdf": "PDF", "docx": "Word document", "pptx": "PowerPoint", "notebook": "Jupyter notebook", "text": "Text file"}


def kind_of(filename: str) -> str | None:
    ext = filename.rsplit(".", 1)[-1].lower()
    return next((k for k, exts in SUPPORTED_EXTENSIONS.items() if ext in exts), None)


# ── text units ───────────────────────────────────────────────────────────────

def read_units(path: str, filename: str) -> tuple[str, tuple[str, ...]]:
    """(unit name, texts) — pages of a PDF, slides, notebook cells, or lines."""
    return _read_units(path, filename, Path(path).stat().st_mtime_ns)


@lru_cache(maxsize=16)
def _read_units(path: str, filename: str, _mtime: int) -> tuple[str, tuple[str, ...]]:
    # cached per file version: paging through a long PDF shouldn't re-parse it each call
    kind = kind_of(filename)
    if kind == "pdf":
        import fitz
        return "page", tuple(page.get_text() for page in fitz.open(path))
    if kind == "pptx":
        from pptx import Presentation
        return "slide", tuple(_slide_text(s) for s in Presentation(path).slides)
    if kind == "notebook":
        return "cell", tuple(_notebook_cells(path))
    if kind == "docx":
        from docx import Document
        return "line", tuple(p.text for p in Document(path).paragraphs)
    if kind == "excel":
        import pandas as pd
        lines = []
        for sheet, df in pd.read_excel(path, sheet_name=None).items():
            lines += [f"### Sheet: {sheet}"] + df.to_csv(index=False).splitlines()
        return "line", tuple(lines)
    if kind in ("csv", "text"):
        return "line", tuple(_read_text(path).splitlines())
    raise ValueError(f"'{filename}' has no text to read")


def _read_text(path: str) -> str:
    raw = Path(path).read_bytes()
    for encoding in ("utf-8", "utf-8-sig", "cp1250"):
        try:
            return raw.decode(encoding)
        except UnicodeDecodeError:
            continue
    return raw.decode("latin-1")


def _slide_text(slide) -> str:
    parts = []
    for shape in slide.shapes:
        if shape.has_text_frame and shape.text_frame.text.strip():
            parts.append(shape.text_frame.text)
        if shape.has_table:
            parts.append("\n".join("\t".join(cell.text for cell in row.cells) for row in shape.table.rows))
    if slide.has_notes_slide and slide.notes_slide.notes_text_frame.text.strip():
        parts.append(f"[notes]\n{slide.notes_slide.notes_text_frame.text}")
    return "\n".join(parts)


def _notebook_cells(path: str) -> list[str]:
    import nbformat
    nb = nbformat.read(open(path, encoding="utf-8", errors="replace"), as_version=4)
    cells = []
    for cell in nb.cells:
        text = f"[{cell.cell_type}]\n{cell.source}"
        for out in cell.get("outputs", []) if cell.cell_type == "code" else []:
            if "text" in out:
                text += f"\n[output]\n{''.join(out.text)}"
            elif "data" in out and "text/plain" in out.data:
                text += f"\n[output]\n{out.data['text/plain']}"
        cells.append(text)
    return cells


def render_unit(unit: str, number: int, text: str) -> str:
    if unit == "line":
        return f"{number:>5}| {text}"
    return f"[{unit} {number}]\n{text.strip()}"


def join_units(unit: str, blocks: list[str]) -> str:
    return ("\n" if unit == "line" else "\n\n").join(blocks)


# ── what goes into the message ───────────────────────────────────────────────

def process_file(path: str, filename: str) -> dict:
    kind = kind_of(filename)
    if kind is None:
        ext = filename.rsplit(".", 1)[-1].lower()
        return {"type": "error", "content": f"Unsupported format '.{ext}'. Supported: {', '.join(ALL_SUPPORTED)}", "name": filename}
    try:
        if kind == "image":
            return {"type": "image", "name": filename}
        if kind == "excel":
            return _excel_preview(path, filename)
        if kind == "csv":
            return _csv_preview(path, filename)
        return _document_preview(path, filename, kind)
    except Exception as e:
        return {"type": "error", "content": f"Error reading file: {str(e)}", "name": filename}


def _document_preview(path: str, filename: str, kind: str) -> dict:
    unit, items = read_units(path, filename)
    text = join_units(unit, [render_unit(unit, i, t) for i, t in enumerate(items, 1)])
    if not "".join(items).strip():
        reason = "it has no text layer (looks like a scan) and OCR isn't available" if kind == "pdf" else "it has no text"
        return {"type": "error", "content": f"Can't read {filename}: {reason}.", "name": filename}

    size = f"{LABEL[kind]}, {len(items)} {unit}{'' if len(items) == 1 else 's'}, {sum(map(len, items)):,} characters"
    if len(text) <= SMALL_FILE_CHARS:
        return {"type": "text", "content": f"({size} — shown in full)\n{text}", "name": filename}
    return {"type": "text", "name": filename, "content": (
        f"({size} — only the beginning is shown)\n{text[:HEAD_CHARS]}…\n\n"
        f"[Read on with read_file('{filename}', start=…) or find a passage with "
        f"read_file('{filename}', query='key words'). Don't answer about parts you haven't read.]"
    )}


def _excel_preview(path: str, filename: str) -> dict:
    import pandas as pd

    sheets = []
    for sheet_name, df in pd.read_excel(path, sheet_name=None).items():
        sheets.append(
            f"### Sheet: {sheet_name} (shape: {df.shape[0]} rows × {df.shape[1]} cols)\n"
            f"Columns: {', '.join(str(c) for c in df.columns)}\n\n"
            f"First {min(MAX_PREVIEW_ROWS, len(df))} rows:\n{df.head(MAX_PREVIEW_ROWS).to_markdown(index=False)}"
        )
    text = "\n\n".join(sheets)
    text += f"\n\n(Full data in the Python sandbox — pd.read_excel('/workspace/{filename}', sheet_name=...))"
    return {"type": "text", "content": text, "name": filename}


def _csv_preview(path: str, filename: str) -> dict:
    import csv as _csv
    import pandas as pd

    for encoding in ["utf-8", "utf-8-sig", "latin-1", "cp1250"]:
        try:
            content = open(path, encoding=encoding).read()
            try:
                dialect = _csv.Sniffer().sniff(content[:2048], delimiters=",;\t|")
                sep = dialect.delimiter
            except Exception:
                sep = ","
            decimal = "."
            for line in content.splitlines()[1:6]:
                if any("," in p for p in line.split(sep)):
                    decimal = ","
                    break

            df = pd.read_csv(path, sep=sep, decimal=decimal, encoding=encoding)
            text = (
                f"[CSV: sep='{sep}', decimal='{decimal}'; shape: {df.shape[0]} rows × {df.shape[1]} cols]\n"
                f"Columns: {', '.join(str(c) for c in df.columns)}\n\n"
                f"First {min(MAX_PREVIEW_ROWS, len(df))} rows:\n{df.head(MAX_PREVIEW_ROWS).to_markdown(index=False)}\n\n"
                f"(Full data in the Python sandbox — "
                f"pd.read_csv('/workspace/{filename}', sep='{sep}', decimal='{decimal}'))"
            )
            return {"type": "text", "content": text, "name": filename}
        except UnicodeDecodeError:
            continue
    return {"type": "error", "content": "Could not read the CSV file.", "name": filename}
